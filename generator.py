from openpyxl import load_workbook
from pathlib import Path
import html


EXCEL_FILE = "colecciones.xlsx"
OUTPUT_FILE = "index.html"


def clean(value):
    """Convierte None en string vacío y escapa HTML."""
    if value is None:
        return ""

    return html.escape(str(value).strip())


def generate_card(row):
    id = clean(row["id"])
    marca = clean(row["marca"])
    producto = clean(row["producto"])
    lugar = clean(row["lugar de adquisición"])
    descripcion = clean(row["descripción"])
    imagen = clean(row["imagen"])

    product_html = (
        f'<p class="card-product">{producto}</p>'
        if producto else ""
    )

    place_html = (
        f'<p class="card-place">{lugar}</p>'
        if lugar else ""
    )

    description_html = (
        f'<p class="card-description">{descripcion}</p>'
        if descripcion else ""
    )

    return f"""
        <article class="card"
            data-brand="{marca}"
            data-place="{lugar}"
        >

            <div class="card-image">
                <img src="assets/tapitas/{id}.png" alt="Tapita {marca}">
            </div>

            <div class="card-info">
                <span class="card-id">#{id}</span>

                <h3>{marca}</h3>

                {product_html}
                {place_html}
                {description_html}
            </div>

        </article>
    """


def read_sheet(sheet):
    """Lee una hoja y devuelve sus filas como diccionarios."""

    rows = sheet.iter_rows(values_only=True)

    headers = next(rows)

    return [
        dict(zip(headers, row))
        for row in rows
        if any(value is not None for value in row)
    ]


def generate_collection_section(sheet):
    collection_name = sheet.title

    rows = read_sheet(sheet)

    cards = "\n".join(
        generate_card(row)
        for row in rows
    )

    places = sorted({
        row.get("lugar de adquisición")
        for row in rows
        if row.get("lugar de adquisición")
    })

    filters_html = """
        <button class="filter-chip active" data-place="">
            Todas
        </button>
    """

    filters_html += "\n".join(
        f"""
        <button class="filter-chip" data-place="{clean(place)}">
            {clean(place)}
        </button>
        """
        for place in places
    )

    section_id = collection_name.lower().replace(" ", "-")

    return f"""
    <section id="{section_id}" class="collection">

        <h2>{collection_name}</h2>

        <p class="collection-count">
            {len(rows)} piezas
        </p>

        <div class="filters">

            <input
                type="search"
                class="search-input"
                placeholder="Buscar en la colección..."
            >

            <div class="filter-chips">
                {filters_html}
            </div>

        </div>

        <div class="grid">
            {cards}
        </div>

    </section>
    """


def generate_collections(workbook):
    return "\n".join(
        generate_collection_section(sheet)
        for sheet in workbook.worksheets
    )


def main():

    workbook = load_workbook(EXCEL_FILE)

    collections_html = generate_collections(workbook)

    template = Path(OUTPUT_FILE).read_text(encoding="utf-8")

    marker = '<!-- COLLECTIONS -->'

    if marker not in template:
        raise ValueError(
            f"No se encontró el marcador {marker} en index.html"
        )

    output = template.replace(
        marker,
        collections_html
    )

    Path(OUTPUT_FILE).write_text(
        output,
        encoding="utf-8"
    )

    print("✓ Catálogo generado correctamente.")


if __name__ == "__main__":
    main()